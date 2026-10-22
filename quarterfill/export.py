from io import BytesIO
import json
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def safe(value):
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def workbook(cells):
    wb = Workbook()
    fin = wb.active
    fin.title = 'Quarterly Financials'
    audit = wb.create_sheet('Calculation Audit')
    sources = wb.create_sheet('SEC Sources')
    quality = wb.create_sheet('Data Quality')
    fin.append(['Company', 'Ticker', 'Metric', 'Fiscal year', 'Quarter', 'Start', 'End', 'Value', 'Unit', 'Status', 'Verified', 'Audit ID'])
    audit.append(['Cell ID', 'Formula / component graph', 'Input IDs', 'Reasoning', 'Warnings'])
    sources.append(['Fact ID', 'Value', 'Unit', 'Concept', 'Label', 'Start', 'End', 'Form', 'Filed', 'Accession', 'SEC document', 'Context', 'Dimensions', 'Excerpt', 'Verified'])
    quality.append(['Cell ID', 'Status', 'Verified', 'Warnings', 'Investigation'])
    indices = {}
    for cell in cells:
        for f in cell.inputs:
            if f.id not in indices:
                sources.append([safe(x) for x in [f.id, float(f.value), f.unit, f.concept, f.label, f.start, f.end,
                    f.form, f.filed, f.accession, f.url, f.context, str(f.dimensions), f.section, f.verified]])
                indices[f.id] = sources.max_row
        value = float(cell.value) if cell.value is not None else None
        if cell.value is not None and cell.operation in ('', '-', '+'):
            refs = [f"'SEC Sources'!B{indices[f.id]}" for f in cell.inputs]
            value = '=' + (f' {cell.operation} '.join(refs) if cell.operation else refs[0])
        fin.append([safe(x) for x in [cell.company, cell.ticker, cell.metric, cell.year, cell.quarter,
                   cell.start, cell.end, None, cell.unit, cell.status, cell.verified, cell.id]])
        fin.cell(fin.max_row, 8, value)
        fin.cell(fin.max_row, 8).number_format = '#,##0.00;[Red](#,##0.00);–'
        audit.append([safe(x) for x in [cell.id, cell.formula, ', '.join(f.id for f in cell.inputs),
                     '\n'.join(cell.reasoning), '\n'.join(cell.warnings)]])
        quality.append([cell.id, cell.status, cell.verified, '\n'.join(cell.warnings), json.dumps(cell.investigation)])
    for ws in wb:
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for header in ws[1]:
            header.font = Font(color='FFFFFF', bold=True)
            header.fill = PatternFill('solid', fgColor='18283B')
        for column in ws.columns:
            letter = get_column_letter(column[0].column)
            ws.column_dimensions[letter].width = min(65, max(14, max(len(str(c.value or '')) for c in column) + 2))
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical='top', wrap_text=True)
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def fill_model(content, mappings, cells):
    """Explicit mapping only: {cell ID: (sheet, address)}. Existing formulas are protected."""
    wb = load_workbook(BytesIO(content))
    changes = []
    for c in cells:
        if c.id not in mappings or c.value is None or not c.verified:
            continue
        sheet, address = mappings[c.id]
        target = wb[sheet][address]
        if target.value is not None:
            changes.append((sheet, address, 'Skipped: nonempty'))
            continue
        target.value = float(c.value)
        from openpyxl.comments import Comment
        target.comment = Comment(json.dumps(c.audit(), indent=2), 'QuarterFill')
        changes.append((sheet, address, 'Filled with verified value; audit in comment'))
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue(), changes
