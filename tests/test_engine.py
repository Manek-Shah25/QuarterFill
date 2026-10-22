from dataclasses import replace
from decimal import Decimal
from io import BytesIO
import pytest
from openpyxl import load_workbook, Workbook
from quarterfill.models import Fact
from quarterfill.concepts import BY_NAME, Metric, search
from quarterfill.engine import reconstruct, fiscal_periods, derive
from quarterfill.export import workbook, fill_model
from quarterfill.sec import parse_inline, filing_url, verify, SECError


def fact(value, start='2024-01-01', end='2024-03-31', **kw):
    args = dict(cik='1', concept='us-gaap:ShareBasedCompensation', label='SBC',
        value=Decimal(str(value)), unit='USD', start=start, end=end,
        accession='0000000001-24-000001', form='10-Q', filed='2024-05-01',
        url='https://www.sec.gov/Archives/edgar/data/1/000000000124000001/report.htm', verified=True)
    args.update(kw)
    return Fact(**args)


PERIODS = ('2024-01-01', {1:'2024-03-31',2:'2024-06-30',3:'2024-09-30',4:'2024-12-31'})
FLOWS = [fact(100), fact(260,end='2024-06-30'),fact(450,end='2024-09-30'),
         fact(700,end='2024-12-31',form='10-K')]


def cell(facts, q, metric='Stock-based compensation'):
    return reconstruct('Test fixture', 'TEST', BY_NAME[metric], facts, 2024, q, PERIODS)


@pytest.mark.parametrize('q,expected', [(1,100),(2,160),(3,190),(4,250),(0,700)])
def test_flow_reconstruction(q, expected):
    c = cell(FLOWS,q)
    assert c.value == expected
    assert c.verified
    assert c.audit()['input_ids']
    assert c.formula.endswith('USD')


def test_direct_preferred():
    direct = fact(160,start='2024-04-01',end='2024-06-30')
    assert cell(FLOWS+[direct],2).status == 'Directly reported'


def test_balance_never_subtracted():
    f = fact(50,start=None,concept='us-gaap:CashAndCashEquivalentsAtCarryingValue')
    assert cell([f],1,'Cash').value == 50
    assert cell([f],2,'Cash').value is None


@pytest.mark.parametrize('name,unit,tag', [('Diluted EPS','USD/shares','EarningsPerShareDiluted'),
    ('Diluted weighted-average shares','shares','WeightedAverageNumberOfDilutedSharesOutstanding')])
def test_nonadditive(name,unit,tag):
    fs = [replace(f, concept='us-gaap:'+tag, unit=unit) for f in FLOWS]
    assert cell(fs,2,name).value is None


def test_ratio_never_differenced():
    m = Metric('Margin','Income','nonadditive',('Margin',),unit='pure')
    fs = [replace(f,concept='us-gaap:Margin',unit='pure') for f in FLOWS]
    assert reconstruct('T','T',m,fs,2024,2,PERIODS).value is None


def test_restated_conflict_preserves_versions():
    amended = replace(FLOWS[1],value=Decimal(270),form='10-Q/A',accession='0000000001-24-000002')
    c = cell(FLOWS+[amended],2)
    assert c.value is None and c.status == 'Ambiguous or conflicting'
    assert amended in c.inputs


def test_missing_is_not_zero():
    assert cell([],1).value is None
    assert cell([fact(0)],1).status == 'Explicitly zero'


def test_dimension_exclusion():
    assert cell([fact(100, dimensions=('SegmentA',))],1).value is None


def test_incompatible_unit_and_tag():
    assert cell([replace(f,unit='EUR') for f in FLOWS],2).value is None
    different = replace(FLOWS[0],concept='us-gaap:PaymentsToAcquirePropertyPlantAndEquipment')
    assert cell([different,FLOWS[1]],2).value is None


def test_noncalendar_53_week():
    annual = fact(100,start='2022-09-25',end='2023-09-30',form='10-K')
    q = fact(20,start='2022-09-25',end='2022-12-31')
    assert fiscal_periods([annual,q],2023) == ('2022-09-25',{1:'2022-12-31',4:'2023-09-30'})


def test_no_fiscal_anchor():
    assert fiscal_periods(FLOWS,2025) is None


def test_negative_cashflow_sign_preserved():
    assert cell([replace(f,value=-f.value) for f in FLOWS],4).value == -250


def test_custom_requires_explicit_concept():
    m = Metric('Custom','Custom','flow',())
    fs = [replace(f,concept='test:Custom') for f in FLOWS]
    assert reconstruct('T','T',m,fs,2024,2,PERIODS,'test:Custom').value == 160


def test_audit_stable_and_complete():
    c = cell(FLOWS,4)
    assert c.id == cell(FLOWS,4).id
    assert all(f.url and f.accession and f.start and f.end for f in c.inputs)
    assert c.audit()['formula'] == '700 - 450 = 250 USD'


def test_unverified_not_verified():
    assert not cell([replace(FLOWS[0],verified=False)],1).verified


def test_excel_formulas_and_injection():
    c = cell(FLOWS,4)
    c.company='=HYPERLINK("evil")'
    wb = load_workbook(BytesIO(workbook([c])),data_only=False)
    assert wb['Quarterly Financials']['H2'].value == "='SEC Sources'!B2 - 'SEC Sources'!B3"
    assert wb['Quarterly Financials']['A2'].data_type == 's'
    assert len(wb.sheetnames) == 4


def test_model_preserves_formulas_and_existing_values():
    wb = Workbook(); wb.active.title='Model'; wb.active['A1']='=1+1'; wb.active['B1']=42
    stream=BytesIO(); wb.save(stream)
    c=cell(FLOWS,1)
    for addr in ('A1','B1','C1'):
        out,changes=fill_model(stream.getvalue(),{c.id:('Model',addr)},[c])
        edited=load_workbook(BytesIO(out))
        assert edited['Model']['A1'].value=='=1+1'
        assert edited['Model']['B1'].value==42
        if addr=='C1':
            assert edited['Model'][addr].value==100 and edited['Model'][addr].comment


def test_filing_url_only_from_metadata():
    with pytest.raises(SECError):
        filing_url(1,{'accessionNumber':'invalid','primaryDocument':'fake.htm'})


def test_inline_scale_extension_and_context():
    html='''<html xmlns:xbrli="x" xmlns:ix="i"><xbrli:context id="c"><xbrli:period>
    <xbrli:startDate>2024-01-01</xbrli:startDate><xbrli:endDate>2024-03-31</xbrli:endDate>
    </xbrli:period></xbrli:context><xbrli:unit id="u"><xbrli:measure>iso4217:USD</xbrli:measure></xbrli:unit>
    <table><tr><td>Custom expense</td><td><ix:nonFraction name="test:Custom" contextRef="c" unitRef="u"
    scale="6" sign="-" decimals="-6" format="ixt:num-dot-decimal">1,200</ix:nonFraction></td></tr></table></html>'''
    parsed,text=parse_inline(html,1,{'accessionNumber':'0000000001-24-000001','form':'10-Q','filingDate':'2024-05-01'},'https://www.sec.gov/real')
    assert parsed[0].value == -1200000000
    assert parsed[0].verified and parsed[0].concept=='test:Custom'
    assert verify(replace(parsed[0],verified=False),parsed).verified


def test_definition_search():
    assert search('CapEx')[0].name=='Capital expenditures'


def test_derived_missing_and_valid():
    a=cell(FLOWS,2); b=cell(FLOWS,1)
    assert derive('T','T','FCF',[a,b],['-']).value is None  # incompatible quarter ends
    b=replace(a,value=Decimal(20))
    assert derive('T','T','FCF',[a,b],['-']).value==140


def test_empty_unit_no_imputation():
    assert cell([replace(FLOWS[0],unit='')],1).value is None
